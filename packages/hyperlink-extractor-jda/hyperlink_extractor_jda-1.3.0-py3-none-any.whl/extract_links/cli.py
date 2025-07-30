#!/usr/bin/env python3
import argparse, csv, json, re
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from odf.opendocument import load as odf_load
from odf.text import P, A

def extract_from_html(path):
    with open(path, encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'html.parser')
    results = []
    for a in soup.find_all('a', href=True):
        text = ' '.join(a.stripped_strings)
        text = re.sub(r'\s+', ' ', text).strip()
        href = a['href'].strip()
        if not href: continue
        label = text if text and text != href else '[No text]'
        results.append((label, href))
    return results

def extract_text_recursive(node):
    result = []
    for child in node.childNodes:
        if child.nodeType == child.TEXT_NODE:
            result.append(child.data)
        else:
            result.append(extract_text_recursive(child))
    return ''.join(result)

def extract_from_odt(path):
    doc = odf_load(str(path))
    results = []
    for p in doc.getElementsByType(P):
        for a in p.getElementsByType(A):
            text = extract_text_recursive(a).strip()
            href = a.getAttribute('href')
            if not href: continue
            label = text if text and text != href else '[No text]'
            results.append((label, href))
    return results

def extract_from_docx(path):
    doc = Document(path)
    rels = doc.part.rels
    results = []
    for paragraph in doc.paragraphs:
        p = paragraph._p
        for hl in p.findall(".//w:hyperlink", namespaces=p.nsmap):
            rId = hl.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            if rId is None: continue
            url = rels[rId].target_ref.strip()
            text = ''.join([n.text for n in hl.findall(".//w:t", namespaces=p.nsmap) if n.text]).strip()
            label = text if text and text != url else '[No text]'
            results.append((label, url))
    return results

def extract_from_pdf(path):
    import fitz  # PyMuPDF
    import unicodedata
    results = []
    doc = fitz.open(str(path))
    for page in doc:
        words = page.get_text("words")  # list of (x0, y0, x1, y1, word, block_no, line_no, word_no)
        for link in page.get_links():
            uri = link.get('uri')
            rect = link.get('from')
            if not uri:
                continue

            r = fitz.Rect(rect)
            # Inflate manually to catch text slightly outside the box
            r.x0 -= 2
            r.y0 -= 2
            r.x1 += 2
            r.y1 += 2

            # Collect words that intersect the link rectangle
            nearby_words = [w for w in words if fitz.Rect(w[:4]).intersects(r)]
            # Sort words by vertical, then horizontal position
            nearby_words.sort(key=lambda w: (round(w[1], 1), w[0]))

            text = ' '.join(w[4] for w in nearby_words).strip()

            # Clean & normalize
            text = unicodedata.normalize("NFKC", text)
            text = re.sub(r'\s+', ' ', text)

            label = text if text and text != uri else '[No text]'
            results.append((label, uri.strip()))
    return results

def dedupe(results):
    seen = set()
    deduped = []
    for label, url in results:
        if url not in seen:
            deduped.append((label, url))
            seen.add(url)
    return deduped

def sort_results(results):
    return sorted(results, key=lambda x: x[0].lower())

def write_outputs(results, output_prefix, fmt='txt'):
    if fmt == 'txt':
        with open(f"{output_prefix}.txt", 'w', encoding='utf-8') as f:
            for text, url in results:
                f.write(f"{text} — {url}\n")
    elif fmt == 'csv':
        with open(f"{output_prefix}.csv", 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(['Text', 'URL'])
            writer.writerows(results)
    elif fmt == 'md':
        with open(f"{output_prefix}.md", 'w', encoding='utf-8') as f:
            for text, url in results:
                f.write(f"[{text}]({url})\n")
    elif fmt == 'json':
        with open(f"{output_prefix}.json", 'w', encoding='utf-8') as f:
            json.dump([{'text': t, 'url': u} for t,u in results], f, indent=2, ensure_ascii=False)
    elif fmt == 'xlsx':
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Text", "URL"])
        for text, url in results:
            ws.append([text, url])
        wb.save(f"{output_prefix}.xlsx")

def main():
    parser = argparse.ArgumentParser(description="Extract hyperlinks from .docx, .odt, .html, or .pdf")
    parser.add_argument("input", help="Input file")
    parser.add_argument("-o", "--output", help="Output prefix", default="output")
    parser.add_argument("--dedupe", action='store_true', help="Remove duplicate URLs")
    parser.add_argument("--sort", action='store_true', help="Sort by label")
    parser.add_argument("--format", choices=['txt','csv','md','json','xlsx'], default='txt', help="Output format")
    args = parser.parse_args()

    input_path = Path(args.input)
    ext = input_path.suffix.lower()

    if ext == ".html":
        results = extract_from_html(input_path)
    elif ext == ".odt":
        results = extract_from_odt(input_path)
    elif ext == ".docx":
        results = extract_from_docx(input_path)
    elif ext == ".pdf":
        results = extract_from_pdf(input_path)
    else:
        raise ValueError("Unsupported file type. Use .docx, .odt, .html, or .pdf")

    if args.dedupe:
        results = dedupe(results)
    if args.sort:
        results = sort_results(results)

    write_outputs(results, args.output, args.format)

    print(f"✅ Extracted {len(results)} links. Output: {args.output}.{args.format}")

def run():
    main()

if __name__ == "__main__":
    main()
